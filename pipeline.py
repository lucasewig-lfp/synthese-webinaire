"""Étapes du traitement : transcription (mlx-whisper), synthèse (Claude Code), export Word."""

import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).parent
WHISPER_CLI = Path(sys.executable).parent / "mlx_whisper"
WHISPER_MODEL = "mlx-community/whisper-large-v3-turbo"
CLAUDE_CANDIDATES = [Path.home() / ".local/bin/claude", Path.home() / ".claude/local/claude",
                     Path("/opt/homebrew/bin/claude"), Path("/usr/local/bin/claude")]


# ---------------------------------------------------------------- transcription

def transcribe(media_path: Path, out_dir: Path, language: str, on_progress) -> str:
    """Extrait l'audio puis le transcrit en local avec mlx-whisper. Renvoie le texte."""
    if not shutil.which("ffmpeg"):
        raise RuntimeError("ffmpeg introuvable. Lance installer.command.")

    wav = out_dir / "audio.wav"
    on_progress("Extraction de l'audio…", 2)
    subprocess.run(
        ["ffmpeg", "-y", "-i", str(media_path), "-vn", "-ac", "1", "-ar", "16000", "-c:a", "pcm_s16le", str(wav)],
        check=True, capture_output=True,
    )

    on_progress("Transcription en cours… (le 1er lancement télécharge le modèle, 1,6 Go)", 5)
    cmd = [str(WHISPER_CLI), str(wav), "--model", WHISPER_MODEL, "--output-dir", str(out_dir),
           "--output-name", "transcription", "--output-format", "txt", "--verbose", "False",
           "--condition-on-previous-text", "False"]
    if language != "auto":
        cmd += ["--language", language]
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)

    # la barre de progression (tqdm) se réécrit avec \r : on lit par morceaux
    tail, buf = [], b""
    while chunk := proc.stdout.read(256):
        buf += chunk
        *parts, buf = re.split(rb"[\r\n]", buf)
        for part in parts:
            line = part.decode(errors="replace")
            tail = (tail + [line])[-20:]
            if m := re.search(r"(\d+)%\|", line):
                pct = int(m.group(1))
                on_progress(f"Transcription en cours… {pct} %", 5 + int(pct * 0.75))
    if proc.wait() != 0:
        raise RuntimeError("Échec de la transcription :\n" + "\n".join(tail))
    wav.unlink(missing_ok=True)

    text = (out_dir / "transcription.txt").read_text(encoding="utf-8")
    # whisper découpe en segments d'une ligne : on recolle en paragraphes lisibles
    return re.sub(r"\s*\n\s*", " ", text).strip()


# ---------------------------------------------------------------- synthèse

SYSTEM_PROMPT = """Tu rédiges des synthèses écrites de webinaires à partir de leur transcription automatique.

On te fournit des exemples de synthèses déjà produites (PDF). Ils définissent le modèle à suivre : reprends leur structure (sections, ordre, intitulés), leur longueur approximative, leur ton et leur niveau de détail. N'en recopie jamais le contenu : seul le webinaire fourni sert de source.

La transcription est automatique : corrige silencieusement les erreurs évidentes (noms propres, termes techniques, chiffres mal reconnus) quand le contexte les rend certaines. N'invente aucune information absente de la transcription ; si une rubrique du modèle n'a pas d'équivalent dans le webinaire, indique-le brièvement.

Réponds uniquement avec le document final, en Markdown simple : titres avec #, ## et ###, listes avec - ou 1., **gras**, *italique*, tableaux Markdown si le modèle en contient. Pas de préambule ni de commentaire autour du document."""


def find_claude() -> str | None:
    found = shutil.which("claude")
    if found:
        return found
    return next((str(p) for p in CLAUDE_CANDIDATES if p.exists()), None)


def _claude_env() -> dict:
    # on force l'abonnement Claude : une éventuelle clé API dans l'environnement serait prioritaire
    env = {k: v for k, v in os.environ.items() if k not in ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN")}
    env["PATH"] = f"{Path.home() / '.local/bin'}:{env.get('PATH', '')}"
    return env


def claude_status() -> dict:
    """Claude Code est-il installé et connecté à un compte Claude ?"""
    claude = find_claude()
    if not claude:
        return {"installed": False, "logged_in": False}
    try:
        out = subprocess.run([claude, "auth", "status"], capture_output=True, text=True, timeout=20, env=_claude_env())
        logged_in = bool(json.loads(out.stdout).get("loggedIn"))
    except (subprocess.SubprocessError, ValueError):
        logged_in = False
    return {"installed": True, "logged_in": logged_in}


def summarize(transcript: str, example_pdfs: list[Path], user_prompt: str, work_dir: Path) -> str:
    """Fait rédiger la synthèse par Claude Code, avec l'abonnement Claude de l'utilisateur."""
    claude = find_claude()
    if not claude:
        raise RuntimeError("Claude Code n'est pas installé. Suis l'étape « Connexion à Claude » en haut de la page.")

    # Claude lit lui-même les PDF d'exemple, copiés dans le dossier du travail
    ex_dir = work_dir / "exemples"
    ex_dir.mkdir(exist_ok=True)
    names = []
    for pdf in example_pdfs:
        shutil.copy(pdf, ex_dir / pdf.name)
        names.append(f"exemples/{pdf.name}")

    if names:
        prompt = ("Commence par lire en entier, avec l'outil Read, ces exemples de synthèse : "
                  + ", ".join(names) + ".\nPuis rédige la synthèse du webinaire ci-dessous en suivant leur modèle.")
    else:
        prompt = "Rédige la synthèse du webinaire ci-dessous."
    if user_prompt.strip():
        prompt += f"\n\nConsignes supplémentaires :\n{user_prompt.strip()}"
    prompt += f"\n\n<transcription>\n{transcript}\n</transcription>"

    try:
        proc = subprocess.run(
            [claude, "-p", "--output-format", "json", "--model", "opus", "--effort", "high",
             "--system-prompt", SYSTEM_PROMPT, "--tools", "Read", "--allowedTools", "Read",
             "--strict-mcp-config", "--no-session-persistence"],
            input=prompt, capture_output=True, text=True, cwd=work_dir, env=_claude_env(), timeout=30 * 60,
        )
    except subprocess.TimeoutExpired:
        raise RuntimeError("Claude n'a pas répondu en 30 minutes. Réessaie avec « Régénérer ».")

    try:
        result = json.loads(proc.stdout)
        text = (result.get("result") or "").strip()
        failed = result.get("is_error") or not text
    except ValueError:
        text, failed = (proc.stdout + proc.stderr).strip()[-500:], True
    if failed:
        if "login" in text.lower() or "logged in" in text.lower():
            raise RuntimeError("Claude Code n'est pas connecté à ton compte. Suis l'étape « Connexion à Claude » en haut de la page.")
        if "limit" in text.lower():
            raise RuntimeError(f"Quota de ton abonnement Claude atteint : {text}. Réessaie plus tard avec « Régénérer ».")
        raise RuntimeError(f"Erreur de Claude : {text or 'réponse vide'}")
    # au cas où le document serait entouré d'un bloc ```markdown
    return re.sub(r"^```(?:markdown|md)?\s*\n(.*)\n```$", r"\1", text, flags=re.S)


# ---------------------------------------------------------------- export Word

INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _add_inline(paragraph, text: str):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _table(doc, rows: list[str]):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not cells:
        return
    width = max(len(r) for r in cells)
    table = doc.add_table(rows=len(cells), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, value in enumerate(row):
            p = table.cell(i, j).paragraphs[0]
            _add_inline(p, value)
            if i == 0:
                for run in p.runs:
                    run.bold = True


def markdown_to_docx(markdown: str, out_path: Path):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            _table(doc, rows)
            continue

        if not stripped:
            pass
        elif m := re.match(r"^(#{1,4})\s+(.*)", stripped):
            level = len(m.group(1))
            heading = doc.add_heading(level=0 if level == 1 else level - 1)
            _add_inline(heading, m.group(2))
        elif re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", stripped):
            doc.add_paragraph("_" * 40)
        elif m := re.match(r"^(\s*)[-*•]\s+(.*)", line):
            depth = min(len(m.group(1).expandtabs(4)) // 2, 2)
            style = "List Bullet" if depth == 0 else f"List Bullet {depth + 1}"
            _add_inline(doc.add_paragraph(style=style), m.group(2))
        elif m := re.match(r"^(\s*)\d+[.)]\s+(.*)", line):
            depth = min(len(m.group(1).expandtabs(4)) // 2, 2)
            style = "List Number" if depth == 0 else f"List Number {depth + 1}"
            _add_inline(doc.add_paragraph(style=style), m.group(2))
        elif stripped.startswith(">"):
            _add_inline(doc.add_paragraph(style="Quote"), stripped.lstrip("> "))
        else:
            _add_inline(doc.add_paragraph(), stripped)
        i += 1

    doc.save(out_path)
